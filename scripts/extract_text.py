from pathlib import Path
import json
import re
import unicodedata

import fitz  # PyMuPDF
from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[1]

RAW_DIR = PROJECT_ROOT / "data" / "raw"
TEXT_OUTPUT_DIR = PROJECT_ROOT / "data" / "processed" / "text"
REPORT_OUTPUT_PATH = PROJECT_ROOT / "data" / "processed" / "text_extraction_report.json"


SUPPORTED_EXTENSIONS = {".docx", ".pdf"}


def normalize_text(text: str) -> str:
    """
    Normalize extracted text without destroying Arabic content.

    Purpose:
    - Standardize Unicode characters.
    - Remove excessive spaces.
    - Remove excessive empty lines.
    - Keep paragraph boundaries readable.
    """

    text = unicodedata.normalize("NFKC", text)

    # Normalize Windows/Mac line endings.
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Remove excessive spaces and tabs inside each line.
    lines = []
    for line in text.split("\n"):
        clean_line = re.sub(r"[ \t]+", " ", line).strip()
        lines.append(clean_line)

    text = "\n".join(lines)

    # Remove too many blank lines.
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def extract_docx_text(file_path: Path) -> str:
    """
    Extract text from a DOCX file.

    Purpose:
    - Read normal paragraphs.
    - Read tables too, because legal documents often contain amendments,
      article references, or structured information inside tables.
    """

    document = Document(file_path)
    parts: list[str] = []

    # Extract normal paragraphs.
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    # Extract tables.
    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"\n[Table {table_index}]")

        for row in table.rows:
            cells = []
            for cell in row.cells:
                cell_text = normalize_text(cell.text)
                if cell_text:
                    cells.append(cell_text)

            if cells:
                parts.append(" | ".join(cells))

    return normalize_text("\n".join(parts))


def extract_pdf_text(file_path: Path) -> str:
    """
    Extract text from a PDF file.

    Purpose:
    - Read page by page.
    - Add page markers, which can help debugging later if the PDF extraction
      has formatting problems.
    """

    parts: list[str] = []

    with fitz.open(file_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            page_text = page.get_text("text")
            page_text = normalize_text(page_text)

            if page_text:
                parts.append(f"\n[Page {page_number}]\n{page_text}")

    return normalize_text("\n".join(parts))


def extract_text_from_file(file_path: Path) -> str:
    """
    Choose the correct extractor based on file extension.
    """

    extension = file_path.suffix.lower()

    if extension == ".docx":
        return extract_docx_text(file_path)

    if extension == ".pdf":
        return extract_pdf_text(file_path)

    raise ValueError(f"Unsupported file type: {file_path}")


def build_output_filename(file_path: Path) -> str:
    """
    Convert source file name into a TXT file name.

    Example:
    Labor_Law_EN.pdf -> labor_law_en.txt
    social_insurance_law_ar.docx -> social_insurance_law_ar.txt
    """

    name = file_path.stem.lower()
    name = re.sub(r"[^a-zA-Z0-9_\u0600-\u06FF]+", "_", name)
    name = re.sub(r"_+", "_", name).strip("_")

    return f"{name}.txt"


def find_source_files() -> list[Path]:
    """
    Find all supported raw files inside data/raw.

    Purpose:
    - Keep the script flexible.
    - It will work even if your filenames are slightly different.
    """

    files = []

    for file_path in RAW_DIR.rglob("*"):
        if file_path.is_file() and file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
            # Ignore temporary Word files.
            if file_path.name.startswith("~$"):
                continue

            files.append(file_path)

    return sorted(files)


def main() -> None:
    TEXT_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    source_files = find_source_files()

    if not source_files:
        print("No DOCX or PDF files found inside data/raw.")
        print("Put your files inside:")
        print("  data/raw/arabic/")
        print("  data/raw/english/")
        return

    report = []

    print("Starting text extraction...")
    print("=" * 70)

    for source_file in source_files:
        relative_source = source_file.relative_to(PROJECT_ROOT)

        try:
            extracted_text = extract_text_from_file(source_file)

            output_filename = build_output_filename(source_file)
            output_path = TEXT_OUTPUT_DIR / output_filename

            output_path.write_text(extracted_text, encoding="utf-8")

            item = {
                "source_file": str(relative_source),
                "output_file": str(output_path.relative_to(PROJECT_ROOT)),
                "file_type": source_file.suffix.lower(),
                "characters": len(extracted_text),
                "lines": len(extracted_text.splitlines()),
                "status": "success",
            }

            report.append(item)

            print(f"Extracted: {relative_source}")
            print(f"Output:    {output_path.relative_to(PROJECT_ROOT)}")
            print(f"Chars:     {len(extracted_text):,}")
            print("-" * 70)

        except Exception as error:
            item = {
                "source_file": str(relative_source),
                "output_file": None,
                "file_type": source_file.suffix.lower(),
                "characters": 0,
                "lines": 0,
                "status": "failed",
                "error": str(error),
            }

            report.append(item)

            print(f"Failed: {relative_source}")
            print(f"Error:  {error}")
            print("-" * 70)

    REPORT_OUTPUT_PATH.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    successful = sum(1 for item in report if item["status"] == "success")
    failed = sum(1 for item in report if item["status"] == "failed")

    print("=" * 70)
    print("Text extraction completed")
    print(f"Successful files: {successful}")
    print(f"Failed files:     {failed}")
    print(f"Report saved to:  {REPORT_OUTPUT_PATH.relative_to(PROJECT_ROOT)}")


if __name__ == "__main__":
    main()